"""Negative slide schema paths and blocks the live deck does not use."""

from __future__ import annotations

from pathlib import Path

import pytest
import yaml
from pptx import Presentation

from deck.build_deck import build
from deck.schema import (
    ContentSlide,
    DiagramBlock,
    Section,
    format_value,
    load_deck,
    parse_run,
    validate_components,
    visible_text,
)

ROOT = Path(__file__).resolve().parents[1]
NOTES = "<!-- slide: open.one -->\n1. Cover - 0:00-0:10\nHello\n"


def _dump(document: object) -> str:
    return yaml.safe_dump(document, sort_keys=False)


def _load(
    tmp_path: Path,
    document: object,
    notes: str = NOTES,
    **kwargs: object,
) -> object:
    content = tmp_path / "content"
    content.mkdir()
    (content / "00.yaml").write_text(_dump(document), encoding="utf-8")
    path = tmp_path / "notes.md"
    path.write_text(notes, encoding="utf-8")
    return load_deck(content, path, root=ROOT, **kwargs)  # type: ignore[arg-type]


def _slide(**extra: object) -> dict[str, object]:
    base: dict[str, object] = {
        "id": "open.one",
        "kind": "content",
        "eyebrow": "EYE",
        "title": "Title",
        "deck": "Deck",
        "layout": "full",
        "main": [],
    }
    base.update(extra)
    return base


def _doc(
    slides: list[object] | None = None,
    *,
    section: object | None = None,
    version: object = 1,
) -> dict[str, object]:
    return {
        "schema_version": version,
        "section": section
        or {
            "id": "open",
            "index": "00",
            "title": "Opening",
            "minutes": [0, 1],
        },
        "slides": [_slide()] if slides is None else slides,
    }


def test_visible_lineage_image_and_scalars(tmp_path: Path) -> None:
    document = _doc(
        [
            _slide(
                main=[
                    {
                        "type": "lineage",
                        "ids": ["titan-text", "bedrock"],
                    },
                    {
                        "type": "metrics",
                        "items": [
                            {"value": True, "label": "Yes"},
                            {"value": False, "label": "No"},
                            {"value": 1200, "label": "Count"},
                            {"value": 1.5, "label": "Ratio"},
                        ],
                    },
                    {
                        "type": "image",
                        "path": "deck/assets/logo_dark.png",
                        "caption": "Logo",
                    },
                    {"type": "image", "path": "deck/assets/logo_white.png"},
                    {
                        "type": "text",
                        "text": "Plain {lineage.titan-text.name}",
                    },
                    {
                        "type": "table",
                        "source": {
                            "lineage": {
                                "filter": {"id": ["titan-text"]},
                                "fields": [
                                    "name",
                                    "former_names",
                                    "successors",
                                ],
                            }
                        },
                    },
                    {
                        "type": "timeline",
                        "ids": ["titan-text"],
                    },
                ]
            )
        ]
    )
    deck = _load(tmp_path, document)
    rows = deck.slides[0].main[0].rows  # type: ignore[attr-defined]
    titan = rows[0]
    assert titan["after"] != "-"
    assert " | " in titan["dates"]
    assert "-" in rows[1]["after"]
    assert format_value(True) == "yes"
    assert format_value(False) == "no"
    output = tmp_path / "image.pptx"
    build(output, deck, draft=True)
    assert output.is_file()


@pytest.mark.parametrize(
    ("document", "message"),
    [
        ("not-a-mapping", "Expected a mapping"),
        ({"schema_version": True, "section": {}, "slides": []}, "schema"),
        ({"schema_version": 2, "section": {}, "slides": []}, "schema"),
        ({"slides": []}, "Missing fields"),
        (
            {
                "schema_version": 1,
                "section": {
                    "id": "open",
                    "index": "00",
                    "title": "Opening",
                    "minutes": [1],
                },
                "slides": [],
            },
            "minutes",
        ),
        (
            {
                "schema_version": 1,
                "section": {
                    "id": "open",
                    "index": "00",
                    "title": "Opening",
                    "minutes": [2, 1],
                },
                "slides": [],
            },
            "increase",
        ),
    ],
)
def test_deck_envelope_errors(
    tmp_path: Path, document: object, message: str
) -> None:
    with pytest.raises(ValueError, match=message):
        _load(tmp_path, document)


def test_slide_and_block_errors(tmp_path: Path) -> None:
    errors = [
        (_doc([_slide(kind="nope")]), "Unknown value"),
        (_doc([_slide(id="other.one")]), "must start"),
        (
            _doc(
                [_slide(layout="full", rail=[{"type": "text", "text": "x"}])]
            ),
            "rail",
        ),
        (_doc([_slide(run=["not-a-command"])]), "awsai-demo"),
        (_doc([_slide(run=["uv run awsai-demo list --help"])]), "awsai-demo"),
        (_doc([_slide(main=[{"type": "nope"}])]), "Unknown block"),
        (
            _doc(
                [_slide(main=[{"type": "text", "text": "x", "bold": "yes"}])]
            ),
            "boolean",
        ),
        (
            _doc(
                [
                    _slide(
                        main=[
                            {
                                "type": "bullets",
                                "items": [{"text": "x", "level": -1}],
                            }
                        ]
                    )
                ]
            ),
            "level",
        ),
        (
            _doc(
                [
                    _slide(
                        main=[
                            {
                                "type": "bullets",
                                "items": [{"text": "x", "level": True}],
                            }
                        ]
                    )
                ]
            ),
            "level",
        ),
        (
            _doc(
                [
                    _slide(
                        main=[
                            {
                                "type": "metrics",
                                "items": [{"value": None, "label": "x"}],
                            }
                        ]
                    )
                ]
            ),
            "scalars",
        ),
        (_doc([_slide(main=[{"type": "table"}])]), "exactly one"),
        (
            _doc(
                [
                    _slide(
                        main=[
                            {
                                "type": "table",
                                "rows": [["a"]],
                                "source": {"facts": "x"},
                            }
                        ]
                    )
                ]
            ),
            "exactly one",
        ),
        (
            _doc(
                [
                    _slide(
                        main=[
                            {
                                "type": "table",
                                "columns": ["A", "B"],
                                "rows": [["only"]],
                            }
                        ]
                    )
                ]
            ),
            "match",
        ),
        (
            _doc(
                [
                    _slide(
                        main=[
                            {
                                "type": "table",
                                "columns": ["A"],
                                "rows": [["a"]],
                                "widths": [1, 2],
                            }
                        ]
                    )
                ]
            ),
            "widths",
        ),
        (
            _doc([_slide(main=[{"type": "lineage", "ids": ["missing"]}])]),
            "Unknown lineage",
        ),
        (
            _doc([_slide(main=[{"type": "image", "path": "README.md"}])]),
            "deck/assets",
        ),
        (
            _doc(
                [
                    _slide(
                        main=[
                            {
                                "type": "image",
                                "path": "deck/assets/missing.png",
                            }
                        ]
                    )
                ]
            ),
            "deck/assets",
        ),
        (_doc([_slide(main=[{"type": "code", "ref": "nope"}])]), "Invalid"),
        (
            _doc(
                [
                    _slide(
                        main=[
                            {
                                "type": "diagram",
                                "nodes": [
                                    {
                                        "id": "a",
                                        "label": "A",
                                        "x": 0,
                                        "y": 0,
                                        "w": 1,
                                        "h": 1,
                                    },
                                    {
                                        "id": "a",
                                        "label": "B",
                                        "x": 2,
                                        "y": 0,
                                        "w": 1,
                                        "h": 1,
                                    },
                                ],
                                "edges": [],
                            }
                        ]
                    )
                ]
            ),
            "Duplicate",
        ),
        (
            _doc([_slide(main=[{"h": True, "type": "text", "text": "x"}])]),
            "number",
        ),
        (
            _doc([_slide(main=[{"h": 0, "type": "text", "text": "x"}])]),
            "positive",
        ),
        (
            _doc([_slide(main=[{"h": -1, "type": "text", "text": "x"}])]),
            "positive",
        ),
        (_doc(slides=[]), "must exist"),
        (_doc([_slide(), _slide()]), "unique"),
        ({1: "x"}, "string keys"),
        (
            _doc([_slide(main=[{"type": "links", "groups": []}])]),
            "at least one group",
        ),
        (
            _doc(
                [
                    _slide(
                        main=[
                            {
                                "type": "links",
                                "groups": [{"title": "Docs", "items": []}],
                            }
                        ]
                    )
                ]
            ),
            "at least one link",
        ),
        (
            _doc(
                [
                    _slide(
                        main=[
                            {
                                "type": "links",
                                "groups": [
                                    {
                                        "title": "Docs",
                                        "items": [
                                            {
                                                "label": "Plain",
                                                "url": "http://example.com",
                                            }
                                        ],
                                    }
                                ],
                            }
                        ]
                    )
                ]
            ),
            "https",
        ),
    ]
    for document, message in errors:
        with pytest.raises(ValueError, match=message):
            _load(tmp_path, document)
        _clear(tmp_path)


def _clear(tmp_path: Path) -> None:
    for child in tmp_path.iterdir():
        if child.is_dir():
            for item in child.iterdir():
                item.unlink()
            child.rmdir()
        else:
            child.unlink()


def test_draft_records_diagram_overflow_and_bad_braces(tmp_path: Path) -> None:
    node = {"id": "a", "label": "A", "x": 30, "y": 0, "w": 1, "h": 1}
    diagram = _doc(
        [_slide(main=[{"type": "diagram", "nodes": [node], "edges": []}])]
    )
    with pytest.raises(ValueError, match="outside block"):
        _load(tmp_path, diagram)
    _clear(tmp_path)
    problems: list[str] = []
    _load(tmp_path, diagram, draft=True, problems=problems)
    assert any("outside block" in item for item in problems)
    _clear(tmp_path)
    braced = _doc([_slide(title="Bad {brace")])
    with pytest.raises(ValueError, match="interpolation"):
        _load(tmp_path, braced)
    _clear(tmp_path)
    problems = []
    _load(tmp_path, braced, draft=True, problems=problems)
    assert any("interpolation" in item for item in problems)


def test_components_and_duplicate_sections(tmp_path: Path) -> None:
    deck = _load(tmp_path, _doc())
    path = tmp_path / "components.yaml"
    path.write_text(
        "components:\n  - id: one\n    slides: [open.one]\n",
        encoding="utf-8",
    )
    validate_components(deck, path)  # type: ignore[arg-type]
    path.write_text(
        "components:\n  - id: one\n    slides: [missing.slide]\n",
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="missing slides"):
        validate_components(deck, path)  # type: ignore[arg-type]
    other = tmp_path / "content" / "01.yaml"
    other.write_text(_dump(_doc()), encoding="utf-8")
    with pytest.raises(ValueError, match="unique"):
        load_deck(
            tmp_path / "content",
            tmp_path / "notes.md",
            root=ROOT,
        )


def test_remaining_schema_edges(tmp_path: Path) -> None:
    section = {
        "id": "open",
        "index": "00",
        "title": "Opening",
        "minutes": [0, 1],
    }
    failures: tuple[tuple[object, str], ...] = (
        (_doc([_slide(title="")]), "nonempty"),
        (
            {"schema_version": 1, "section": section, "slides": "nope"},
            "list",
        ),
        (
            _doc([_slide(title="{lineage.titan-text.not_a_field}")]),
            "Unknown lineage field",
        ),
        (_doc([_slide(title="{facts.demo_count}")]), "resolver"),
        (_doc([_slide(main=[{"type": "timeline"}])]), "exactly one"),
        (
            _doc(
                [
                    _slide(
                        main=[
                            {
                                "type": "timeline",
                                "ids": ["titan-text"],
                                "events": [
                                    {"date": "2024-01-01", "label": "Start"}
                                ],
                            }
                        ]
                    )
                ]
            ),
            "exactly one",
        ),
        (_doc([_slide(id="open.two")]), "Missing notes"),
    )
    for document, message in failures:
        with pytest.raises(ValueError, match=message):
            _load(tmp_path, document)
        _clear(tmp_path)
    deck = _load(
        tmp_path,
        _doc(
            [
                _slide(
                    main=[
                        {
                            "type": "timeline",
                            "events": [
                                {"date": "2024-02-01", "label": "Later"},
                                {
                                    "date": "2020-01-01",
                                    "label": "Earlier",
                                    "lane": "a",
                                },
                            ],
                        },
                        {
                            "type": "layer_map",
                            "layers": [
                                {
                                    "name": "Core",
                                    "items": [{"label": "Bedrock"}],
                                }
                            ],
                        },
                    ]
                )
            ]
        ),
    )
    events = deck.slides[0].main[0].events  # type: ignore[attr-defined]
    assert events[0]["label"] == "Earlier"
    layers = deck.slides[0].main[1].layers  # type: ignore[attr-defined]
    assert layers[0]["name"] == "Core"
    _clear(tmp_path)
    extra_notes = NOTES + "<!-- slide: open.extra -->\nExtra\n"
    with pytest.raises(ValueError, match="without slides"):
        _load(tmp_path, _doc(), notes=extra_notes)
    _clear(tmp_path)
    notes = NOTES + "<!-- slide: open.two -->\n2. Next - 0:10-0:20\nSecond\n"
    (tmp_path / "notes.md").write_text(notes, encoding="utf-8")
    content = tmp_path / "content"
    content.mkdir()
    (content / "00.yaml").write_text(_dump(_doc()), encoding="utf-8")
    second = _doc([_slide(id="open.two")])
    (content / "01.yaml").write_text(_dump(second), encoding="utf-8")
    with pytest.raises(ValueError, match="Section ids must be unique"):
        load_deck(content, tmp_path / "notes.md", root=ROOT)
    shown = visible_text(
        ContentSlide(
            id="open.one",
            kind="content",
            eyebrow="E",
            title="T",
            deck="D",
            section=Section("open", "00", "Opening", (0, 1)),
            notes="n",
            main=(
                DiagramBlock(
                    type="diagram",
                    nodes=({"label": "A", "score": 1},),
                    edges=(),
                ),
            ),
        )
    )
    assert "A" in shown
    assert "1" not in shown


def test_missing_table_fact_uses_an_explicit_placeholder(
    tmp_path: Path,
) -> None:
    def resolve(path: str) -> str:
        return f"[MISSING: {path}: absent]"

    deck = _load(
        tmp_path,
        _doc(
            [
                _slide(
                    main=[
                        {
                            "type": "table",
                            "columns": ["Concern", "Lane"],
                            "source": {"facts": "offline.decision.matrix"},
                        }
                    ]
                )
            ]
        ),
        resolve=resolve,
    )
    rows = deck.slides[0].main[0].rows  # type: ignore[attr-defined]
    assert rows == (("[MISSING: facts.offline.decision.matrix: absent]", ""),)


def test_identifier_metrics_skip_thousands_separators(
    tmp_path: Path,
) -> None:
    def resolve(path: str) -> int:
        assert path.startswith("facts.")
        return 25740

    deck = _load(
        tmp_path,
        _doc(
            [
                _slide(
                    main=[
                        {
                            "type": "metrics",
                            "items": [
                                {
                                    "value": (
                                        "{facts.offline.decision.pause_pid}"
                                    ),
                                    "label": "pid",
                                },
                                {
                                    "value": "{facts.offline.scenario.total}",
                                    "label": "total",
                                },
                            ],
                        }
                    ]
                )
            ]
        ),
        resolve=resolve,
    )
    items = deck.slides[0].main[0].items  # type: ignore[attr-defined]
    assert items[0]["value"] == "25740"
    assert items[1]["value"] == "25,740"
    assert format_value(25740, kind="identifier") == "25740"
    assert format_value("pid-1", kind="identifier") == "pid-1"


def test_diagram_boundary_does_not_depend_on_slide_copy(
    tmp_path: Path,
) -> None:
    deck = _load(
        tmp_path,
        _doc(
            [
                _slide(
                    main=[
                        {
                            "type": "diagram",
                            "h": 2.0,
                            "nodes": [
                                {
                                    "id": "a",
                                    "label": "A",
                                    "x": 0.2,
                                    "y": 0.4,
                                    "w": 1.6,
                                    "h": 0.7,
                                },
                                {
                                    "id": "b",
                                    "label": "B",
                                    "x": 2.4,
                                    "y": 0.4,
                                    "w": 1.6,
                                    "h": 0.7,
                                },
                            ],
                            "edges": [
                                {
                                    "from": "a",
                                    "to": "b",
                                    "label": "go",
                                    "style": "dashed",
                                }
                            ],
                            "boundaries": [
                                {"label": "Group", "nodes": ["a", "b"]}
                            ],
                        }
                    ]
                )
            ]
        ),
    )
    boundary = deck.slides[0].main[0].boundaries[0]  # type: ignore[attr-defined]
    assert boundary["label"] == "Group"
    output = tmp_path / "boundary.pptx"
    build(output, deck, draft=True)
    shown = [
        shape.text_frame.text
        for shape in Presentation(str(output)).slides[0].shapes
        if shape.has_text_frame
    ]
    assert "Group" in shown


def test_parse_run_accepts_a_real_command() -> None:
    assert parse_run("uv run awsai-demo list")[0] == "list"
    with pytest.raises(ValueError, match="non-null"):
        format_value({"nope": 1})
