"""Draft deck build and the checks that do not depend on slide copy."""

from __future__ import annotations

import json
import runpy
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu, Pt

from deck import components, theme
from deck.build_deck import build, check_body_layout, main
from deck.checks import check_layout, check_text_fit, estimated_height
from deck.notes import attach_notes, load_notes, notes_time_window
from deck.schema import ContentSlide, Deck, Section, TextBlock
from deck.snippets import extract_snippet

ROOT = Path(__file__).resolve().parents[1]


def _missing_fact_root(root: Path) -> Path:
    """Deck whose only measured value is still uncollected."""
    deck = root / "deck"
    (deck / "content").mkdir(parents=True)
    (root / "data").mkdir()
    (deck / "facts_contract.yaml").write_text(
        "schema_version: 1\n"
        "offline:\n"
        "  demo_count: {source: list, how: count, type: integer}\n"
        "live: {}\n",
        encoding="utf-8",
    )
    lane = {
        "fingerprint": None,
        "timestamp": None,
        "sources": {},
        "missing": {},
    }
    offline = dict(lane)
    offline["missing"] = {"demo_count": "not collected"}
    (deck / "facts.json").write_text(
        json.dumps(
            {
                "schema_version": 1,
                "offline": {"demo_count": None},
                "live": {},
                "metadata": {"offline": offline, "live": lane},
            }
        ),
        encoding="utf-8",
    )
    (deck / "content" / "00.yaml").write_text(
        "schema_version: 1\n"
        "section: {id: open, index: '00', title: Opening, minutes: [0, 1]}\n"
        "slides:\n"
        "  - id: open.title\n"
        "    kind: title\n"
        "    eyebrow: WORKSHOP\n"
        '    title: "Count {facts.offline.demo_count}"\n'
        "    deck: Deck\n",
        encoding="utf-8",
    )
    (deck / "speaker_notes.md").write_text(
        "<!-- slide: open.title -->\n1. Cover - 0:00-0:45\nHello\n",
        encoding="utf-8",
    )
    (root / "data" / "lineage.yaml").write_text(
        "schema_version: 1\n"
        "verified: '2026-10-05'\n"
        "records:\n"
        "  - id: bedrock\n"
        "    name: Amazon Bedrock\n"
        "    former_names: []\n"
        "    entity_type: service\n"
        "    scope: Models\n"
        "    official_status: Available\n"
        "    normalized_status: active\n"
        "    launched: '2023-09-28'\n"
        "    announced: null\n"
        "    effective: null\n"
        "    end_of_support: null\n"
        "    successors: []\n"
        "    relation: none\n"
        "    sources: [https://docs.aws.amazon.com/bedrock/]\n"
        "    note: null\n",
        encoding="utf-8",
    )
    return root


def test_draft_build_keeps_review_problems(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    output = tmp_path / "aws_ai_stack.pptx"
    assert main(["--draft", "--output", str(output)]) == 0
    assert output.is_file()
    assert main(["--draft", "--output", str(output)], root=tmp_path) == 1
    capsys.readouterr()
    draft = tmp_path / "draft.pptx"
    fixture = _missing_fact_root(tmp_path / "fixture")
    assert main(["--draft", "--output", str(draft)], root=fixture) == 0
    assert draft.is_file()
    assert (
        "DRAFT: offline.demo_count: not collected" in capsys.readouterr().err
    )


def test_checks_report_grid_headline_and_height(tmp_path: Path) -> None:
    del tmp_path
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(theme.SLIDE_W), Emu(theme.SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(
        str(theme.ASSETS / "logo_dark.png"),
        Emu(0),
        Emu(0),
        Emu(theme.SLIDE_W),
        Emu(theme.SLIDE_H),
    )
    theme.box(
        slide,
        "",
        x=theme.MARGIN_L,
        y=theme.Y_BODY,
        w=theme.inches(2),
        h=theme.inches(0.4),
        name="chrome-blank",
    )
    slide.shapes.add_picture(
        str(theme.ASSETS / "logo_dark.png"),
        Emu(theme.MARGIN_L),
        Emu(theme.Y_BODY),
        Emu(100000),
        Emu(100000),
    )
    theme.box(
        slide,
        "WWWWWWWWWWWWWWWWWWWW",
        x=0,
        y=0,
        w=theme.inches(1),
        h=theme.inches(0.4),
        size=40,
        name="title",
    )
    theme.box(
        slide,
        "low",
        x=theme.MARGIN_L,
        y=theme.SLIDE_H - theme.inches(0.2),
        w=theme.inches(2),
        h=theme.inches(1),
    )
    theme.box(
        slide,
        "right",
        x=theme.SLIDE_W - theme.inches(0.2),
        y=theme.Y_BODY,
        w=theme.inches(1),
        h=theme.inches(0.4),
    )
    theme.box(
        slide,
        "short",
        x=theme.MARGIN_L,
        y=theme.Y_BODY,
        w=theme.inches(6),
        h=theme.inches(1),
        name="title",
    )
    theme.box(
        slide,
        "word " * 80,
        x=theme.MARGIN_L,
        y=theme.Y_BODY,
        w=theme.inches(1),
        h=theme.inches(0.2),
    )
    raw = slide.shapes.add_textbox(
        Emu(100), Emu(100), Emu(2000000), Emu(200000)
    )
    raw.text_frame.paragraphs[0].text = "Hi"
    extra = raw.text_frame.add_paragraph()
    extra.line_spacing = 1.4
    extra.space_after = Pt(2)
    run = extra.add_run()
    run.text = "Wrapped"
    run.font.size = Pt(14)
    empty = raw.text_frame.add_paragraph()
    empty.add_run().text = ""
    assert estimated_height(raw) > 0
    layout = check_layout(prs)
    fit = check_text_fit(prs)
    assert any("into the footer" in item for item in layout)
    assert any("off the right edge" in item for item in layout)
    assert any("sits above the grid" in item for item in layout)
    assert any("sits left of the grid" in item for item in layout)
    assert any("too wide for one line" in item for item in fit)
    assert any("more height" in item for item in fit)


def test_body_layout_reserves_notes_and_commands() -> None:
    section = Section("open", "00", "Opening", (0, 1))

    def slide(note: str, run: tuple[str, ...], height: float) -> ContentSlide:
        return ContentSlide(
            id="open.one",
            kind="content",
            eyebrow="E",
            title="T",
            deck="D",
            section=section,
            notes="notes",
            main=(TextBlock(type="text", text="Hi", h=height),),
            run=run,
            source_note=note,
        )

    plain = check_body_layout(Deck((section,), (slide("", (), 0.2),)))
    noted = check_body_layout(Deck((section,), (slide("source", (), 20),)))
    commanded = check_body_layout(
        Deck((section,), (slide("", ("awsai-demo list",), 20),))
    )
    assert plain == []
    assert noted and commanded


def test_clean_non_draft_build_skips_the_banner(tmp_path: Path) -> None:
    output = tmp_path / "clean.pptx"
    assert build(output, Deck((), ())) == []
    assert output.is_file()


def test_build_deck_module_rejects_unknown_arguments(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(sys, "argv", ["build_deck.py", "--nope"])
    monkeypatch.delitem(sys.modules, "deck.build_deck", raising=False)
    with pytest.raises(SystemExit):
        runpy.run_module("deck.build_deck", run_name="__main__")


def test_non_draft_build_rejects_known_problems(tmp_path: Path) -> None:
    section = Section("open", "00", "Opening", (0, 1))
    content = ContentSlide(
        id="open.one",
        kind="content",
        eyebrow="E",
        title="T",
        deck="D",
        section=section,
        notes="notes",
    )
    deck = Deck((section,), (content,))
    with pytest.raises(ValueError, match="bad"):
        build(tmp_path / "out.pptx", deck, evidence_problems=["bad"])


def test_notes_and_snippets_reject_bad_markers(tmp_path: Path) -> None:
    assert notes_time_window("1. Cover - 0:00-0:45\nHi") == "0:00-0:45"
    assert notes_time_window("Appendix") == "Not presented (appendix)"
    notes = tmp_path / "notes.md"
    notes.write_text(
        "<!-- slide: open.one -->\n# Heading\nHello\n"
        "<!-- slide: open.one -->\nAgain\n",
        encoding="utf-8",
    )
    try:
        load_notes(notes)
    except ValueError as exc:
        assert "Duplicate" in str(exc)
    notes.write_text("<!-- slide: open.one -->\n\n", encoding="utf-8")
    try:
        load_notes(notes)
    except ValueError as exc:
        assert "Empty" in str(exc)
    frame = type("Frame", (), {"text": "kept"})()
    slide = type(
        "Slide",
        (),
        {"notes_slide": type("Notes", (), {"notes_text_frame": frame})()},
    )()
    attach_notes(slide, "script")
    assert frame.text.startswith("script")
    broken = type(
        "Slide",
        (),
        {"notes_slide": type("Notes", (), {"notes_text_frame": None})()},
    )()
    try:
        attach_notes(broken, "script")
    except ValueError as exc:
        assert "notes" in str(exc)

    def write(name: str, body: str) -> None:
        path = tmp_path / "src" / "awsai_demo" / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(body, encoding="utf-8")

    write(
        "ok.py",
        "# slide: demo\n"
        "def demo() -> None:\n"
        "    return None\n" + ("x" * 72) + "\n# end-slide: demo\n",
    )
    assert "def demo" in extract_snippet("src/awsai_demo/ok.py#demo", tmp_path)
    cases = {
        "bad.py": (
            "# slide: demo\n" + ("x" * 73) + "\n# end-slide: demo\n",
            "72",
        ),
        "empty.py": ("# slide: demo\n\n# end-slide: demo\n", "1-14"),
        "long.py": (
            "# slide: demo\n"
            + "\n".join(["line"] * 15)
            + "\n# end-slide: demo\n",
            "1-14",
        ),
        "nest.py": (
            "# slide: demo\n# slide: inner\n# end-slide: demo\n",
            "Nested",
        ),
        "twice.py": (
            "# slide: demo\nok\n# end-slide: demo\n"
            "# slide: demo\nok\n# end-slide: demo\n",
            "Nested",
        ),
        "swap.py": ("# slide: demo\nok\n# end-slide: other\n", "Unbalanced"),
        "open.py": ("# slide: demo\nvalue = 1\n", "Unclosed"),
    }
    for name, (body, message) in cases.items():
        write(name, body)
        try:
            extract_snippet(f"src/awsai_demo/{name}#demo", tmp_path)
        except ValueError as exc:
            assert message in str(exc)
        else:
            raise AssertionError(name)
    for ref in (
        "src/awsai_demo/ok.py",
        "src/awsai_demo/ok.txt#demo",
        "README.md#demo",
        "src/awsai_demo/ok.py#missing",
    ):
        write("ok.txt", "x")
        try:
            extract_snippet(ref, tmp_path)
        except ValueError as exc:
            assert "Invalid" in str(exc) or "Missing" in str(exc)
        else:
            raise AssertionError(ref)


def test_code_caption_sits_below_the_panel_and_is_measured() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(theme.SLIDE_W), Emu(theme.SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    allocated = theme.inches(3.6)
    components.draw(
        slide,
        {
            "type": "code",
            "h": 3.6,
            "code": "\n".join(f"line {index}" for index in range(12)),
            "caption": "Evaluated locally with cedarpy.",
        },
        x=theme.MARGIN_L,
        y=theme.Y_BODY,
        w=theme.COL_MAIN_W,
        root=ROOT,
    )
    panel = next(shape for shape in slide.shapes if shape.name == "code-panel")
    code = next(shape for shape in slide.shapes if shape.name == "code")
    caption = next(
        shape for shape in slide.shapes if shape.name == "code-caption"
    )
    assert panel.height == allocated - caption.height
    assert caption.top == theme.Y_BODY + panel.height
    assert caption.top + caption.height == theme.Y_BODY + allocated
    assert code.top >= panel.top
    assert int(code.top + code.height) <= int(panel.top + panel.height)
    bare = prs.slides.add_slide(prs.slide_layouts[6])
    components.draw(
        bare,
        {"type": "code", "h": 2.0, "code": "ok"},
        x=theme.MARGIN_L,
        y=theme.Y_BODY,
        w=theme.COL_MAIN_W,
        root=ROOT,
    )
    assert not any(shape.name == "code-caption" for shape in bare.shapes)
    bare_panel = next(
        shape for shape in bare.shapes if shape.name == "code-panel"
    )
    assert bare_panel.height == theme.inches(2.0)
    tight = prs.slides.add_slide(prs.slide_layouts[6])
    components.draw(
        tight,
        {
            "type": "code",
            "h": 0.8,
            "code": "\n".join(["return value"] * 12),
            "caption": "This caption is part of the measured block.",
        },
        x=theme.MARGIN_L,
        y=theme.Y_BODY,
        w=theme.inches(4),
        root=ROOT,
    )
    assert any("more height" in item for item in check_text_fit(prs))


def test_lineage_without_successors_draws_name_and_status_only() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    components.lineage(
        slide,
        [
            {
                "before": "Amazon Bedrock",
                "relation": "none",
                "after": "-",
                "status": "active",
                "dates": "2023-09-28",
            },
            {
                "before": "Amazon Titan Text models",
                "relation": "successor",
                "after": "Amazon Nova (first generation)",
                "status": "end_of_life",
                "dates": "2023-09-28 | 2026-07",
            },
        ],
        x=theme.inches(1),
        y=theme.inches(1),
        w=theme.inches(12),
        h=theme.inches(2.4),
    )
    texts = [
        shape.text_frame.text for shape in slide.shapes if shape.has_text_frame
    ]
    assert "Amazon Bedrock" in texts
    assert all("none" not in text for text in texts)
    assert sum("→" in text for text in texts) == 1
    assert any(text.startswith("→ successor →") for text in texts)
    assert any(text.startswith("active\n") for text in texts)
