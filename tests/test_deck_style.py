"""The sibling styling: accents, pills, chips, rail panels and code."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from deck import components, theme
from deck.build_deck import build
from deck.diagrams import diagram
from deck.schema import (
    ClosingSlide,
    ContentSlide,
    Deck,
    Section,
    SectionSlide,
    TitleSlide,
)

ROOT = Path(__file__).resolve().parents[1]
P = theme.P


def _slide() -> Any:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(theme.SLIDE_W), Emu(theme.SLIDE_H)
    return prs.slides.add_slide(prs.slide_layouts[6])


def _fills(slide: Any) -> list[str]:
    return [
        str(shape.fill.fore_color.rgb)
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
    ]


def _text(slide: Any, text: str) -> Any:
    return next(
        shape
        for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text == text
    )


def _colour(shape: Any) -> str:
    return str(shape.text_frame.paragraphs[0].runs[0].font.color.rgb)


def test_code_lines_colour_python_and_keep_every_character() -> None:
    code = "\n".join(
        [
            'x = """open',
            "inside # not a comment",
            'close""" + f(y)  # done',
            "@tool",
            "def run() -> None:",
            "    return 's'",
            "",
            "end",
        ]
    )
    lines = components.code_lines(code)
    joined = ["".join(text for text, _ in runs) for runs in lines]
    assert joined == code.split("\n")
    colours = {
        text: style.get("colour") for runs in lines for text, style in runs
    }
    assert colours['"""open'] == P.yellow
    assert colours["inside # not a comment"] == P.yellow
    assert colours['close"""'] == P.yellow
    assert colours["# done"] == P.code_dim
    assert colours["@tool"] == P.teal
    assert colours["def"] == colours["None"] == colours["return"] == P.sky
    assert colours["'s'"] == P.yellow
    assert colours["end"] is None
    assert lines[6] == [("", {})]


def test_diagram_colours_nodes_by_role_and_kind() -> None:
    slide = _slide()
    nodes = [
        {"id": "a", "label": "Entry", "x": 0, "y": 0, "w": 1, "h": 1},
        {"id": "b", "label": "Step", "x": 2, "y": 0, "w": 1, "h": 1},
        {"id": "c", "label": "Exit", "x": 4, "y": 0, "w": 1, "h": 1},
        {"id": "g", "label": "Guard", "kind": "guard"},
        {"id": "s", "label": "Store", "kind": "store", "w": 2},
        {"id": "e", "label": "External", "kind": "external"},
        {"id": "n", "label": "Note", "kind": "note"},
    ]
    for index, node in enumerate(nodes[3:], 3):
        node.update({"x": 2 * index, "y": 0, "h": 1})
        node.setdefault("w", 1)
    edges = [{"from": "a", "to": "b"}, {"from": "b", "to": "c"}]
    diagram(slide, {"nodes": nodes, "edges": edges}, x=0, y=0, w=0, h=0)
    assert _fills(slide) == [P.blue, P.teal, P.violet, P.amber, P.wash, P.card]
    assert _colour(_text(slide, "Entry")) == P.on_dark
    assert _colour(_text(slide, "Step")) == P.ink_soft
    assert _colour(_text(slide, "Note")) == P.muted
    lid = round(theme.inches(1) * 0.15)
    assert _text(slide, "Store").top == lid + theme.inches(0.04)
    assert _text(slide, "External").top == theme.inches(0.14)


def test_matrix_chips_share_one_colour_per_value() -> None:
    slide = _slide()
    components.draw(
        slide,
        {
            "type": "matrix",
            "h": 2.0,
            "columns": ["Concern", "A", "B"],
            "rows": [["x", "app", "aws"], ["y", "aws", "app"]],
        },
        x=theme.MARGIN_L,
        y=theme.Y_BODY,
        w=theme.inches(10),
        root=ROOT,
    )
    chips = [fill for fill in _fills(slide) if fill != P.rule]
    assert chips == [P.coral, P.blue, P.blue, P.coral]
    assert _text(slide, "x").text_frame.paragraphs[0].runs[0].font.bold


def test_schedule_table_draws_a_minute_bar_and_coloured_rows() -> None:
    slide = _slide()
    block: dict[str, Any] = {
        "type": "table",
        "h": 3.0,
        "columns": ["When", "Segment", "Command"],
        "rows": [["0-2", "Open", "awsai-demo list"], ["2-5", "Close", "-"]],
    }
    components.draw(
        slide,
        block,
        x=theme.MARGIN_L,
        y=theme.Y_BODY,
        w=theme.inches(10),
        root=ROOT,
        schedule={"0-2": P.violet},
    )
    fills = _fills(slide)
    assert fills[:2] == [P.violet, P.blue]
    assert _text(slide, "5 min") and _text(slide, "0 min")
    assert _colour(_text(slide, "Close")) == P.on_dark
    command = _text(slide, "awsai-demo list").text_frame.paragraphs[0]
    assert command.runs[0].font.name == theme.FONT_CODE
    plain = _slide()
    block["rows"] = [["Day 1", "Open", "awsai-demo list"]]
    components.draw(
        plain,
        block,
        x=theme.MARGIN_L,
        y=theme.Y_BODY,
        w=theme.inches(10),
        root=ROOT,
    )
    assert not any(
        shape.has_text_frame and shape.text_frame.text.endswith(" min")
        for shape in plain.shapes
    )


def test_prose_uses_rail_panels_and_sized_callout_bars() -> None:
    slide = _slide()
    area: dict[str, Any] = {
        "x": theme.MARGIN_L,
        "w": theme.inches(8),
        "root": ROOT,
    }
    components.draw(
        slide,
        {"type": "callout", "tone": "warning", "text": "Short.", "h": 2.0},
        y=theme.Y_BODY,
        **area,
    )
    bar = slide.shapes[0]
    assert str(bar.fill.fore_color.rgb) == P.amber
    assert bar.height < theme.inches(2.0)
    rail = _slide()
    components.draw(
        rail,
        {"type": "bullets", "items": ["One", {"text": "Two", "level": 1}]},
        y=theme.Y_BODY,
        dark=True,
        **area,
    )
    components.draw(
        rail,
        {"type": "callout", "tone": "info", "text": "Rail"},
        y=theme.Y_BODY + theme.inches(3),
        dark=True,
        **area,
    )
    components.draw(
        rail,
        {"type": "source_note", "text": "Source"},
        y=theme.Y_BODY + theme.inches(5),
        dark=True,
        **area,
    )
    assert _fills(rail) == [P.panel_deep, P.panel_deep, P.blue]
    assert _text(rail, "• One\n  • Two")
    assert _colour(_text(rail, "Rail")) == P.on_dark_soft
    assert _colour(_text(rail, "Source")) == P.muted


def test_build_draws_cover_pills_divider_and_command_pill(
    tmp_path: Path,
) -> None:
    section = Section("open", "00", "Opening", (0, 1))
    common: dict[str, Any] = {"section": section, "notes": "Say this."}
    slides = (
        TitleSlide(
            id="open.title",
            kind="title",
            eyebrow="W",
            title="T",
            deck="Alpha | Beta",
            **common,
        ),
        SectionSlide(
            id="open.section",
            kind="section",
            eyebrow="Part",
            title="Part one",
            deck="Why",
            bullets=("First",),
            **common,
        ),
        ContentSlide(
            id="open.one",
            kind="content",
            eyebrow="E",
            title="T",
            deck="D",
            run=("awsai-demo list", "awsai-demo doctor"),
            **common,
        ),
        ClosingSlide(
            id="open.close",
            kind="closing",
            eyebrow="Q",
            title="Thanks",
            deck="Bye",
            **common,
        ),
    )
    output = tmp_path / "style.pptx"
    assert build(output, Deck((section,), slides)) == []
    prs = Presentation(str(output))
    cover, divider, content, closing = prs.slides
    assert _text(cover, "Alpha") and _text(cover, "Beta")
    assert _colour(_text(divider, "PART")) == P.coral
    assert _fills(divider)[0] == P.coral
    pill = next(
        shape for shape in content.shapes if shape.name == "chrome-run"
    )
    assert pill.text_frame.text == (
        "▶  uv run awsai-demo list   ·   uv run awsai-demo doctor"
    )
    assert _colour(_text(content, "E")) == P.coral
    assert _text(closing, "Bye")
