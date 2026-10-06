"""Build the AWS workshop from typed content and recorded evidence."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import TYPE_CHECKING

# Direct and module execution use the same package.
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Emu  # noqa: E402

from deck import components  # noqa: E402
from deck import theme as T  # noqa: E402
from deck.checks import check_layout, check_text_fit  # noqa: E402
from deck.notes import attach_notes  # noqa: E402
from deck.schema import block_data, load_deck  # noqa: E402

if TYPE_CHECKING:
    from pptx.presentation import Presentation as PresentationType
    from pptx.slide import Slide as PptxSlide

    from deck.schema import Deck, Slide

DEFAULT_OUTPUT = ROOT / "aws_ai_stack.pptx"
DEFAULT_DATE = "2026-10-05"
DEFAULT_AUTHOR = "Mykola Murha"
DEFAULT_ROLE = "Data and AI Engineer  ·  DataArt"
# Cover pill outlines, in the siblings' order of accents.
PILL_COLOURS = (T.P.teal, T.P.sky, T.P.teal, T.P.yellow, T.P.coral)


def _cover(
    slide: PptxSlide, content: Slide, *, date: str, author: str, role: str
) -> None:
    """Draw the cover and closing: eyebrow, title, pills or deck."""
    T.box(
        slide,
        content.eyebrow.upper(),
        x=T.MARGIN_L,
        y=T.inches(2.9),
        w=T.inches(16),
        h=T.inches(0.4),
        size=15,
        colour=T.P.yellow,
        bold=True,
    )
    T.box(
        slide,
        content.title,
        x=T.MARGIN_L,
        y=T.inches(3.45),
        w=T.inches(16),
        h=T.inches(1.3),
        size=60,
        colour=T.P.on_dark,
        bold=True,
    )
    if content.kind == "closing":
        T.box(
            slide,
            content.deck,
            x=T.MARGIN_L,
            y=T.inches(5.1),
            w=T.inches(15),
            h=T.inches(1.2),
            size=19,
            colour=T.P.on_dark_soft,
        )
        return
    for index, label in enumerate(content.deck.split("|")):
        left = T.MARGIN_L + index * T.inches(3.05)
        T.panel(
            slide,
            x=left,
            y=T.inches(5.1),
            w=T.inches(2.9),
            h=T.inches(0.62),
            fill=T.P.on_dark,
            line=PILL_COLOURS[index % len(PILL_COLOURS)],
            radius=T.PILL,
            alpha=14000,
            line_w=1.5,
        )
        T.box(
            slide,
            label.strip(),
            x=left,
            y=T.inches(5.1),
            w=T.inches(2.9),
            h=T.inches(0.62),
            size=15,
            colour=T.P.on_dark,
            bold=True,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )
    for text, top, size, colour in (
        (author, 6.3, 26, T.P.on_dark),
        (role, 6.95, 18, T.P.on_dark_soft),
        (date, 8.0, 15, T.P.on_dark_faint),
    ):
        T.box(
            slide,
            text,
            x=T.MARGIN_L,
            y=T.inches(top),
            w=T.inches(14),
            h=T.inches(0.55),
            size=size,
            colour=colour,
        )


def _divider(slide: PptxSlide, content: Slide, accent: str) -> None:
    """Draw the siblings' section divider and its contents panel."""
    for step in range(4):
        T.panel(
            slide,
            x=T.MARGIN_L + step * 420624,
            y=2359152,
            w=274320,
            h=274320,
            fill=accent if step == 0 else T.P.on_dark,
            line=None,
            kind=MSO_SHAPE.RECTANGLE,
            alpha=None if step == 0 else 30000,
        )
    for text, top, height, size, colour in (
        (content.section.index, 2798064, 1965960, 128, T.P.on_dark),
        (content.eyebrow.upper(), 4956048, 365760, 15, accent),
        (content.title, 5321808, 1920240, 44, T.P.on_dark),
        (content.deck, 7333488, 1188720, 18, T.P.on_dark_soft),
    ):
        T.box(
            slide,
            text,
            x=T.MARGIN_L,
            y=top,
            w=8229600,
            h=height,
            size=size,
            colour=colour,
            bold=size in {128, 15},
        )
    T.panel(
        slide,
        x=10058400,
        y=2798064,
        w=6784848,
        h=4727448,
        fill=T.P.on_dark,
        line=None,
        alpha=10000,
    )
    T.box(
        slide,
        "IN THIS SECTION",
        x=10533888,
        y=3163824,
        w=5833872,
        h=347472,
        size=13,
        colour=T.P.on_dark_faint,
        bold=True,
    )
    for position, item in enumerate(content.bullets):
        top = 3749040 + position * 822960
        T.panel(
            slide,
            x=10533888,
            y=top + 109728,
            w=201168,
            h=201168,
            fill=accent,
            line=None,
            kind=MSO_SHAPE.OVAL,
        )
        T.box(
            slide,
            item,
            x=10954512,
            y=top,
            w=5367528,
            h=786384,
            size=16.5,
            colour=T.P.on_dark_soft,
        )


def render_slide(
    prs: PresentationType,
    content: Slide,
    number: int,
    *,
    root: Path,
    date: str,
    author: str,
    role: str,
    schedule: dict[str, str] | None = None,
) -> None:
    """Render a typed slide and attach its full speaker script."""
    slide = T.chrome(prs, number, kind=content.kind)
    accent = T.SECTION_COLOURS.get(content.section.id, T.P.coral)
    if content.kind == "section":
        _divider(slide, content, T.dark_accent(accent))
    elif content.kind in {"title", "closing"}:
        _cover(slide, content, date=date, author=author, role=role)
    else:
        T.heading(
            slide, content.eyebrow, content.title, content.deck, accent=accent
        )
        main_w = T.CONTENT_W if content.layout == "full" else T.COL_MAIN_W
        for blocks, left, width in (
            (content.main, T.MARGIN_L, main_w),
            (content.rail, T.COL_RAIL_X, T.COL_RAIL_W),
        ):
            top = T.Y_BODY
            for block in blocks:
                top += components.draw(
                    slide,
                    block_data(block),
                    x=left,
                    y=top,
                    w=width,
                    root=root,
                    dark=left == T.COL_RAIL_X,
                    accent=accent,
                    schedule=schedule,
                ) + T.inches(0.12)
        if content.source_note:
            T.box(
                slide,
                content.source_note,
                x=T.MARGIN_L,
                y=T.Y_FOOTER - T.inches(0.52),
                w=T.CONTENT_W,
                h=T.inches(0.3),
                size=10,
                colour=T.P.muted,
            )
        if content.run:
            T.run_strip(slide, content.run)
    attach_notes(slide, content.notes)


def build(
    output: Path,
    content: Deck,
    *,
    root: Path = ROOT,
    draft: bool = False,
    date: str = DEFAULT_DATE,
    author: str = DEFAULT_AUTHOR,
    role: str = DEFAULT_ROLE,
    evidence_problems: list[str] | None = None,
) -> list[str]:
    """Save a validated deck or a clearly marked inspection draft."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(T.SLIDE_W), Emu(T.SLIDE_H)
    # Agenda rows find their section's accent by its minute range.
    schedule = {
        "-".join(f"{minute:g}" for minute in section.minutes): (
            T.SECTION_COLOURS.get(section.id, T.P.blue)
        )
        for section in content.sections
        if section.minutes
    }
    for number, slide in enumerate(content.slides, 1):
        render_slide(
            prs,
            slide,
            number,
            root=root,
            date=date,
            author=author,
            role=role,
            schedule=schedule,
        )
    problems = list(evidence_problems or ())
    problems.extend(check_layout(prs))
    problems.extend(check_text_fit(prs))
    problems.extend(check_body_layout(content))
    if problems and not draft:
        message = "deck validation failed:\n" + "\n".join(problems)
        raise ValueError(message)
    if draft:
        prs.core_properties.subject = "DRAFT - not for presentation"
        for slide in prs.slides:
            T.box(
                slide,
                "DRAFT - review required",
                x=T.SLIDE_W - T.MARGIN_L - T.inches(5.3),
                y=T.inches(0.2),
                w=T.inches(5.3),
                h=T.inches(0.4),
                size=11,
                colour=T.P.coral,
                name="chrome-draft",
                align=PP_ALIGN.RIGHT,
            )
    prs.core_properties.author = author
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    return problems


def check_body_layout(content: Deck) -> list[str]:
    """Detect a stack colliding with source notes or the run strip."""
    problems: list[str] = []
    for number, slide in enumerate(content.slides, 1):
        reserve = 0.95 if slide.source_note else 0.58 if slide.run else 0.1
        bottom = T.Y_FOOTER - T.inches(reserve)
        for name, blocks in (("main", slide.main), ("rail", slide.rail)):
            top = T.Y_BODY
            for block in blocks:
                top += T.inches(
                    block.h or components.DEFAULT_HEIGHTS[block.type]
                )
                if top > bottom:
                    problems.append(
                        f"slide {number} ({slide.id}): {name} {block.type} "
                        f"overlaps the bottom reserved area by "
                        f"{(top - bottom) / 914400:.2f}in"
                    )
                top += T.inches(0.12)
    return problems


def main(argv: list[str] | None = None, *, root: Path = ROOT) -> int:
    """Load local evidence and build without network calls."""
    from deck.facts import load_facts

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--date", default=DEFAULT_DATE)
    parser.add_argument("--author", default=DEFAULT_AUTHOR)
    parser.add_argument("--role", default=DEFAULT_ROLE)
    parser.add_argument("--draft", action="store_true")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args(argv)
    try:
        facts = load_facts(root / "deck/facts.json", root=root)
        problems = facts.problems
        content = load_deck(
            root / "deck/content",
            root / "deck/speaker_notes.md",
            root=root,
            resolve=lambda path: facts.resolve(path, draft=args.draft),
            draft=args.draft,
            problems=problems,
        )
        problems = build(
            args.output,
            content,
            root=root,
            draft=args.draft,
            date=args.date,
            author=args.author,
            role=args.role,
            evidence_problems=problems,
        )
    except (ValueError, OSError) as error:
        print(f"deck build failed: {error}", file=sys.stderr)
        return 1
    for problem in problems:
        print(f"DRAFT: {problem}", file=sys.stderr)
    print(f"{len(content.slides)} slides -> {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
