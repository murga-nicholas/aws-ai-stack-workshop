"""Sibling layout, text-fit and fixed single-line headline checks."""

from __future__ import annotations

from typing import TYPE_CHECKING, Any

from deck import theme as T

if TYPE_CHECKING:
    from pptx.presentation import Presentation

SAFE_TOP = 457200
SAFE_BOTTOM = T.Y_FOOTER - 91440
SAFE_RIGHT = T.SLIDE_W - 457200
CHAR_WIDTH_RATIO = 0.52
FIT_TOLERANCE = 1.12
LINE_HEIGHT = 1.16
MIN_OVERRUN = 91440
ARIAL_WIDTHS = dict(
    zip(
        "abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ ,.;:-",
        (
            556,
            556,
            500,
            556,
            556,
            278,
            556,
            556,
            222,
            222,
            500,
            222,
            833,
            556,
            556,
            556,
            556,
            333,
            500,
            278,
            556,
            500,
            722,
            500,
            500,
            500,
            667,
            667,
            722,
            722,
            667,
            611,
            778,
            722,
            278,
            500,
            667,
            556,
            833,
            722,
            778,
            667,
            778,
            722,
            667,
            611,
            722,
            667,
            944,
            667,
            667,
            611,
            278,
            278,
            278,
            278,
            278,
            333,
        ),
        strict=True,
    )
)


def describe(shape: Any) -> str:
    """Identify a shape by visible words or by its stable shape name."""
    text = shape.text_frame.text if shape.has_text_frame else ""
    return repr(" ".join(text.split())[:55]) if text else str(shape.name)


def check_layout(prs: Presentation) -> list[str]:
    """Report every shape outside the sibling's safe body grid."""
    problems: list[str] = []
    for number, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if (shape.width, shape.height) == (
                T.SLIDE_W,
                T.SLIDE_H,
            ) or shape.name.startswith("chrome-"):
                continue
            bottom, right = shape.top + shape.height, shape.left + shape.width
            label = f"slide {number}: {describe(shape)}"
            if bottom > SAFE_BOTTOM:
                over = (bottom - SAFE_BOTTOM) / 914400
                problems.append(f"{label} runs {over:.2f}in into the footer")
            if right > SAFE_RIGHT:
                over = (right - SAFE_RIGHT) / 914400
                problems.append(
                    f"{label} runs {over:.2f}in off the right edge"
                )
            if shape.top < SAFE_TOP:
                problems.append(f"{label} sits above the grid")
            if shape.left < T.inches(0.5):
                problems.append(f"{label} sits left of the grid")
    return problems


def estimated_height(shape: Any) -> float:
    """Estimate wrapped text using the sibling's conservative model."""
    width = int(shape.width)
    total = 0.0
    for para in shape.text_frame.paragraphs:
        runs = [run for run in para.runs if run.text]
        if not runs:
            continue
        size = max(
            (run.font.size.pt for run in runs if run.font.size), default=14.0
        )
        per_line = max(int(width / (size * CHAR_WIDTH_RATIO * 12700)), 1)
        lines = max(1, -(-sum(len(run.text) for run in runs) // per_line))
        spacing = para.line_spacing
        spacing = spacing if isinstance(spacing, float) else 1.2
        total += lines * size * spacing * LINE_HEIGHT * 12700
        if para.space_after is not None:
            total += para.space_after.pt * 12700
    return total


def _code_caption(shapes: list[Any], code: Any) -> Any | None:
    """Return the caption drawn directly under a code panel."""
    captions = [shape for shape in shapes if shape.name == "code-caption"]
    if not captions:
        return None
    edge = int(code.top + code.height)
    return min(captions, key=lambda shape: abs(int(shape.top) - edge))


def check_text_fit(prs: Presentation) -> list[str]:
    """Report clipped body text and headlines wider than one line."""
    problems: list[str] = []
    for number, slide in enumerate(prs.slides, 1):
        shapes = list(slide.shapes)
        for shape in shapes:
            if (
                not shape.has_text_frame
                or shape.name.startswith("chrome-")
                or shape.name == "code-caption"
                or not shape.text_frame.text.strip()
            ):
                continue
            if shape.name == "title":
                text = shape.text_frame.text.strip()
                size = shape.text_frame.paragraphs[0].runs[0].font.size.pt
                width = sum(ARIAL_WIDTHS.get(char, 556) for char in text)
                over = (width / 1000 * size * 12700 - shape.width) / 914400
                if over > 0:
                    problems.append(
                        f"slide {number}: headline {text!r} is {over:.2f}in "
                        "too wide for one line"
                    )
            else:
                needed = estimated_height(shape)
                overrun = needed - shape.height * FIT_TOLERANCE
                if shape.name == "code":
                    caption = _code_caption(shapes, shape)
                    if caption is not None:
                        caption_over = (
                            estimated_height(caption)
                            - caption.height * FIT_TOLERANCE
                        )
                        overrun = max(
                            overrun, caption_over, overrun + caption_over
                        )
                if overrun > MIN_OVERRUN:
                    problems.append(
                        f"slide {number}: {describe(shape)} needs "
                        f"{overrun / 914400:.2f}in more height than its box"
                    )
    return problems
