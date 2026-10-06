"""DataArt canvas, palette and chrome adapted from the Microsoft deck.

The measured sibling grid, artwork, Arial/Consolas typefaces and fixed
38-point content headlines are preserved; this copy is self-contained.
Colour follows both siblings: every agenda section owns an accent, and
every command that reproduces a slide sits in the footer pill.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING, Any, TypedDict, cast

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Pt

if TYPE_CHECKING:
    from collections.abc import Callable, Sequence

    from pptx.presentation import Presentation
    from pptx.shapes.autoshape import Shape
    from pptx.slide import Slide

SLIDE_W = 18288000
SLIDE_H = 10287000
MARGIN_L = 1439997
CONTENT_W = SLIDE_W - 2 * MARGIN_L
Y_EYEBROW = 658368
Y_TITLE = 1079998
Y_DECK = 2194560
Y_BODY = 3127248
Y_FOOTER = 9428653
COL_MAIN_W = 11338560
COL_RAIL_X = 13121640
COL_RAIL_W = 3721608
FONT_UI = "Arial"
FONT_CODE = "Consolas"
ASSETS = Path(__file__).resolve().parent / "assets"
PILL = 0.5
CARD = 0.0152


@dataclass(frozen=True)
class Palette:
    """The sibling's measured light and dark colour palette."""

    ink: str = "000000"
    ink_soft: str = "23272F"
    body: str = "333333"
    muted: str = "5A6478"
    faint: str = "8A93A6"
    rule: str = "D8DDE8"
    card: str = "FFFFFF"
    wash: str = "F6F8FC"
    panel: str = "0E1A33"
    panel_deep: str = "1B2440"
    code_fg: str = "E6EAF2"
    code_dim: str = "7C8CAE"
    on_dark: str = "FFFFFF"
    on_dark_soft: str = "D9E2F7"
    on_dark_faint: str = "9FB2D8"
    coral: str = "F0503C"
    teal: str = "2BC6BF"
    amber: str = "FFB133"
    blue: str = "3453AD"
    yellow: str = "FFDE55"
    sky: str = "53CFF8"
    violet: str = "70529F"


P = Palette()
STATUS_COLOURS = {
    "active": P.teal,
    "preview": P.blue,
    "maintenance": P.amber,
    "legacy": P.amber,
    "sunset": P.coral,
    "end_of_life": P.coral,
    "renamed": P.faint,
    "moved": P.faint,
}
# One accent per agenda section: its eyebrows, its agenda bar segment
# and its divider share it, so the room always knows where it is.
SECTION_COLOURS = {
    "open": P.coral,
    "lineage": P.blue,
    "bedrock": P.teal,
    "strands": P.violet,
    "agentcore": P.blue,
    "decision": P.amber,
    "data": P.teal,
    "localstack": P.violet,
    "close": P.coral,
    "appendix": P.blue,
}
# Repeating accents for sibling cards with no meaning of their own.
ACCENTS = (P.coral, P.blue, P.teal, P.violet)
# Fills light enough to carry dark text; every other fill takes white.
LIGHT_FILLS = frozenset({P.teal, P.amber, P.yellow, P.sky, P.wash, P.card})


class Style(TypedDict, total=False):
    """Overrides one run of a paragraph's default text style."""

    colour: str
    bold: bool
    font: str
    size: float
    link: str


Run = tuple[str, Style]


def inches(value: float) -> int:
    """Convert the content author's inch measurements to EMU."""
    return round(value * 914400)


def rgb(value: str) -> RGBColor:
    """Construct a typed RGB colour from the palette's hex string."""
    parse = cast("Callable[[str], RGBColor]", RGBColor.from_string)
    return parse(value)


def on_fill(fill: str) -> str:
    """Pick dark or white text so a solid fill stays legible."""
    return P.ink_soft if fill in LIGHT_FILLS else P.on_dark


def dark_accent(colour: str) -> str:
    """Swap accents that vanish on navy artwork for the sibling teal."""
    return colour if colour in LIGHT_FILLS | {P.coral} else P.teal


def rich(
    slide: Slide,
    paragraphs: Sequence[Sequence[Run]],
    *,
    x: int,
    y: int,
    w: int,
    h: int,
    size: float = 14,
    colour: str = P.body,
    bold: bool = False,
    font: str = FONT_UI,
    name: str = "text",
    align: PP_ALIGN = PP_ALIGN.LEFT,
    anchor: MSO_ANCHOR = MSO_ANCHOR.TOP,
    space_after: float = 0,
) -> Shape:
    """Write padded-free paragraphs of individually styled runs."""
    shape = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    shape.name = name
    frame = shape.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    for index, runs in enumerate(paragraphs):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.alignment = align
        para.line_spacing = 1.15
        para.space_after = Pt(space_after)
        for text, style in runs:
            run = para.add_run()
            run.text = text
            run.font.name = style.get("font", font)
            run.font.size = Pt(style.get("size", size))
            run.font.bold = style.get("bold", bold)
            run.font.color.rgb = rgb(style.get("colour", colour))
            if "link" in style:
                # A real hyperlink: Ctrl+click in PowerPoint.
                run.hyperlink.address = style["link"]
    return shape


def box(
    slide: Slide,
    text: str,
    *,
    x: int,
    y: int,
    w: int,
    h: int,
    size: float = 14,
    colour: str = P.body,
    bold: bool = False,
    font: str = FONT_UI,
    name: str = "text",
    align: PP_ALIGN = PP_ALIGN.LEFT,
    anchor: MSO_ANCHOR = MSO_ANCHOR.TOP,
) -> Shape:
    """Write padded-free text, preserving explicit line breaks."""
    return rich(
        slide,
        [[(line, Style())] for line in text.split("\n")],
        x=x,
        y=y,
        w=w,
        h=h,
        size=size,
        colour=colour,
        bold=bold,
        font=font,
        name=name,
        align=align,
        anchor=anchor,
    )


def panel(
    slide: Slide,
    *,
    x: int,
    y: int,
    w: int,
    h: int,
    fill: str = P.wash,
    line: str | None = P.rule,
    kind: MSO_SHAPE = MSO_SHAPE.ROUNDED_RECTANGLE,
    radius: float = CARD,
    alpha: int | None = None,
    line_w: float = 1,
) -> Shape:
    """Draw a standard card, with the sibling's subtle corner radius.

    A ``line`` of None draws no outline, ``radius`` of ``PILL`` rounds
    the ends fully (on a can it is the lid's depth) and ``alpha``
    (OOXML thousandths) lets template artwork show through the fill.
    """
    shape = slide.shapes.add_shape(kind, Emu(x), Emu(y), Emu(w), Emu(h))
    shape.shadow.inherit = False
    if kind in {MSO_SHAPE.ROUNDED_RECTANGLE, MSO_SHAPE.CAN}:
        shape.adjustments[0] = radius
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if alpha is not None:
        element = cast("Any", shape)._element
        for colour in element.spPr.xpath("./a:solidFill/a:srgbClr"):
            node = OxmlElement("a:alpha")
            node.set("val", str(alpha))
            colour.append(node)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_w)
    return shape


def bar(slide: Slide, *, x: int, y: int, w: int, h: int, colour: str) -> Shape:
    """Draw a square-cornered accent bar, rule or marker."""
    return panel(
        slide,
        x=x,
        y=y,
        w=w,
        h=h,
        fill=colour,
        line=None,
        kind=MSO_SHAPE.RECTANGLE,
    )


def chrome(prs: Presentation, number: int, *, kind: str) -> Slide:
    """Add template artwork, logo and fixed footer without content."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    dark = kind in {"title", "section", "closing"}
    artwork = kind if dark else "content"
    slide.shapes.add_picture(
        str(ASSETS / f"bg_{artwork}.jpeg"),
        Emu(0),
        Emu(0),
        Emu(SLIDE_W),
        Emu(SLIDE_H),
    )
    logo = "logo_white.png" if dark else "logo_dark.png"
    dimensions = (
        (1439997, 960120, 2788920, 461772)
        if dark
        else (14695204, 1300130, 2180378, 359999)
    )
    slide.shapes.add_picture(str(ASSETS / logo), *map(Emu, dimensions))
    colour = P.on_dark_faint if dark else P.faint
    box(
        slide,
        "Your Partner for Progress",
        x=MARGIN_L,
        y=Y_FOOTER,
        w=inches(5),
        h=inches(0.45),
        size=12,
        colour=colour,
        name="chrome-footer",
    )
    box(
        slide,
        str(number),
        x=SLIDE_W - MARGIN_L - inches(2),
        y=Y_FOOTER,
        w=inches(2),
        h=inches(0.45),
        size=12,
        colour=colour,
        name="chrome-page-number",
        align=PP_ALIGN.RIGHT,
    )
    return cast("Slide", slide)


def run_strip(slide: Slide, commands: Sequence[str]) -> None:
    """Print the commands that reproduce a slide in the footer pill.

    Every slide with a demo carries one in the same place, as in both
    siblings, so the room learns where to look. The pill is chrome
    between the tagline and the page number; a long list of commands
    shrinks its type rather than wrapping out of the pill.
    """
    text = "   ·   ".join(f"uv run {command}" for command in commands)
    left, right = inches(4.2), SLIDE_W - inches(4.2)
    pad, glyph = 594360, len(text) * 0.58 * 12700
    size = min(14.0, (right - left - pad) / glyph)
    width = min(right - left, round(glyph * size) + pad)
    x, y, h = (SLIDE_W - width) // 2, Y_FOOTER - 64008, 420624
    pill = panel(slide, x=x, y=y, w=width, h=h, fill=P.wash, radius=PILL)
    pill.name = "chrome-run-pill"
    rich(
        slide,
        [[("▶  ", {"colour": P.teal, "font": FONT_UI}), (text, Style())]],
        x=x,
        y=y,
        w=width,
        h=h,
        size=size,
        colour=P.ink_soft,
        bold=True,
        font=FONT_CODE,
        name="chrome-run",
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def heading(
    slide: Slide,
    eyebrow: str,
    title: str,
    deck: str,
    *,
    accent: str = P.coral,
) -> None:
    """Write the sibling's light content heading at fixed type sizes."""
    box(
        slide,
        eyebrow.upper(),
        x=MARGIN_L,
        y=Y_EYEBROW,
        w=CONTENT_W,
        h=310896,
        size=14,
        colour=accent,
        bold=True,
    )
    box(
        slide,
        title,
        x=MARGIN_L,
        y=Y_TITLE,
        w=13209487,
        h=1051560,
        size=38,
        colour=P.ink,
        name="title",
    )
    box(
        slide,
        deck,
        x=MARGIN_L,
        y=Y_DECK,
        w=CONTENT_W,
        h=566928,
        size=19,
        colour=P.muted,
    )
